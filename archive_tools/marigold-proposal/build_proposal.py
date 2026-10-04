import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parents[1] / 'outputs' / 'elec4240-proposal'
OUT.mkdir(parents=True, exist_ok=True)
data = json.loads((ROOT / 'content.json').read_text(encoding='utf-8'))
assert 200 <= len(data['body'].split()) <= 400
assert '\n' not in data['body']
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(1.27)
sec.header_distance = Cm(.6)
sec.footer_distance = Cm(.6)

for name, size, bold in [('Normal', 11, False), ('Title', 14, True), ('Heading 1', 14, True)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(11)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.widow_control = True
    ppr = style._element.find(qn('w:pPr'))
    if ppr is not None:
        borders = ppr.find(qn('w:pBdr'))
        if borders is not None: ppr.remove(borders)

p = doc.add_paragraph('ELEC4240 Course Project Proposal')
p.runs[0].bold = True
p.paragraph_format.keep_with_next = True
p = doc.add_paragraph(data['title'], style='Title')
p.paragraph_format.keep_with_next = True
p = doc.add_paragraph()
for i, (name, sid) in enumerate(data['authors']):
    if i: p.add_run().add_break()
    p.add_run(f'{name} ({sid})')
p.paragraph_format.keep_with_next = True
body = doc.add_paragraph(data['body'])
body.alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.add_paragraph('References', style='Heading 1')
for ref in data['references']:
    p = doc.add_paragraph(ref)
    p.paragraph_format.keep_together = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_after = Pt(0)
r = footer.add_run()
r.font.name = 'Calibri'
r.font.size = Pt(11)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
field_r = OxmlElement('w:r')
field_t = OxmlElement('w:t')
field_t.text = '1'
field_r.append(field_t)
field.append(field_r)
r._r.addnext(field)
props = doc.core_properties
props.title = data['title']
props.subject = 'ELEC4240 Course Project Proposal'
props.author = '; '.join(name for name, _ in data['authors'])
props.keywords = 'Marigold, LoRA, NYU Depth V2, depth estimation'
props.comments = ''
dest = OUT / 'ELEC4240_Marigold_Proposal.docx'
doc.save(dest)
(ROOT / 'authoring_checks.json').write_text(json.dumps({
    'body_word_count': len(data['body'].split()),
    'body_paragraphs': 1,
    'authors': data['authors'],
    'font': 'Calibri', 'body_font_pt': 11, 'title_font_pt': 14,
    'line_spacing': 1.15,
    'margins_cm': {'top': 2.54, 'bottom': 1.27, 'left': 1.27, 'right': 1.27},
    'format_reference_scope': 'Typography only; ELEC4240 single-paragraph scope retained'
}, indent=2), encoding='utf-8')
print(dest)
print('BODY_WORDS', len(data['body'].split()))
